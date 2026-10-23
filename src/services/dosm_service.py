"""Official DOSM quarterly online-offering context, kept separate from skill signals."""

from __future__ import annotations

import io
import json
import os
from datetime import datetime, timezone
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from src.services.storage import connection, initialize_database

SOURCE_URL = "https://www.dosm.gov.my/uploads/release-content/file_20260818075331.xlsx"
SOURCE_PAGE = "https://www.dosm.gov.my/portal-main/release-content/big-data-analytics-job-market-insights-second-quarter-2026"
MAX_BYTES = 8_000_000


def parse_workbook(content: bytes, source_url: str, source_page: str = SOURCE_PAGE) -> dict:
    """Read the occupation table by its headers, comparing the latest quarter YoY."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("DOSM workbook has no active sheet.")
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next(
            (i for i, row in enumerate(rows) if str(row[1]).strip() == "Occupation"), None
        )
        if header_index is None or header_index == 0:
            raise ValueError("DOSM occupation table was not found.")
        year_row, quarter_row = rows[header_index - 1], rows[header_index]
        periods: list[tuple[int, int, str]] = []
        year = None
        for column in range(10, len(quarter_row)):
            if isinstance(year_row[column], int):
                year = year_row[column]
            quarter = quarter_row[column]
            if year and isinstance(quarter, str) and quarter in {"Q1", "Q2", "Q3", "Q4"}:
                total = rows[header_index + 1][column]
                if isinstance(total, (int, float)) and total > 0:
                    periods.append((year, column, quarter))
        if not periods:
            raise ValueError("DOSM workbook has no complete quarterly occupation data.")
        latest_year, latest_column, latest_quarter = periods[-1]
        previous = next(
            ((year, column, quarter) for year, column, quarter in periods
             if year == latest_year - 1 and quarter == latest_quarter), None
        )
        if previous is None:
            raise ValueError("DOSM workbook has no same-quarter previous-year comparison.")
        previous_year, previous_column, _ = previous
        current_total = int(rows[header_index + 1][latest_column])
        previous_total = int(rows[header_index + 1][previous_column])
        groups = []
        for row in rows[header_index + 2:]:
            label = row[1]
            if not isinstance(label, str) or label.startswith("Source:"):
                break
            if label == "Unclassified":
                continue
            current, prior = row[latest_column], row[previous_column]
            if not isinstance(current, (int, float)):
                continue
            current_share = round(current / current_total * 100, 2)
            previous_share = round(prior / previous_total * 100, 2) if isinstance(prior, (int, float)) else None
            shift = round(current_share - previous_share, 2) if previous_share is not None else None
            groups.append({
                "name": label,
                "current_count": int(current),
                "previous_count": int(prior) if previous_share is not None else None,
                "current_share_percent": current_share,
                "previous_share_percent": previous_share,
                "share_point_change": shift,
                "direction": "not_comparable" if shift is None else "higher_share" if shift > 0 else "lower_share" if shift < 0 else "unchanged_share",
            })
        if not groups or sum(group["current_count"] for group in groups) != current_total:
            raise ValueError("DOSM occupation groups do not reconcile with the published total.")
        return {
            "source_name": "Department of Statistics Malaysia (DOSM)",
            "source_url": source_url,
            "source_page": source_page,
            "scope": "Malaysia · online job offerings · major occupation groups",
            "latest_period": f"{latest_quarter} {latest_year}",
            "comparison_period": f"{latest_quarter} {previous_year}",
            "latest_total": current_total,
            "comparison_total": previous_total,
            "groups": groups,
            "limitation": "These are broad occupation groups, not individual roles or skill requirements. Share changes describe this DOSM online-offering series; they do not forecast future demand or prove why a group changed.",
        }
    finally:
        workbook.close()


def get_dosm_dashboard() -> dict:
    initialize_database()
    with connection() as db:
        row = db.execute("SELECT payload, retrieved_at FROM dosm_snapshot WHERE id = 1").fetchone()
    if not row:
        return {"status": "empty", "data": None, "retrieved_at": None}
    return {"status": "available", "data": json.loads(row["payload"]), "retrieved_at": row["retrieved_at"]}


def refresh_dosm() -> dict:
    url = os.getenv("DOSM_WORKBOOK_URL", SOURCE_URL)
    page_url = os.getenv("DOSM_RELEASE_PAGE_URL", SOURCE_PAGE)
    if not url.startswith("https://www.dosm.gov.my/") or not page_url.startswith("https://www.dosm.gov.my/"):
        raise ValueError("DOSM source URLs must use the official DOSM HTTPS domain.")
    request = Request(url, headers={"User-Agent": "Aiversecity/1.0 (source-attributed education prototype)"})
    try:
        with urlopen(request, timeout=25) as response:
            content = response.read(MAX_BYTES + 1)
        if len(content) > MAX_BYTES:
            raise ValueError("DOSM workbook exceeded the download limit.")
        data = parse_workbook(content, url, page_url)
    except (HTTPError, URLError, TimeoutError) as exc:
        raise ValueError("DOSM workbook could not be downloaded.") from exc
    retrieved_at = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    initialize_database()
    with connection() as db:
        db.execute(
            "INSERT INTO dosm_snapshot (id, payload, retrieved_at) VALUES (1, ?, ?) "
            "ON CONFLICT(id) DO UPDATE SET payload = excluded.payload, retrieved_at = excluded.retrieved_at",
            (json.dumps(data), retrieved_at),
        )
    return get_dosm_dashboard()
